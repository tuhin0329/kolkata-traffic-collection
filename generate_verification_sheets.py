import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('all_routes_for_verification.json', 'r', encoding='utf-8') as f:
    routes = json.load(f)

dest_dirs = ['D:/Traffic', 'D:/Traffic/Bus/Data Collection 2', 'D:/Traffic/conclusion']
for d in dest_dirs:
    os.makedirs(d, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Car & Bike Route Links"

# Styles
navy_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
title_font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
sub_font = Font(name="Calibri", size=9, italic=True, color="595959")
regular_font = Font(name="Calibri", size=9)
bold_font = Font(name="Calibri", size=9, bold=True)
link_font = Font(name="Calibri", size=9, color="0563C1", underline="single")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.append(["KOLKATA TRAFFIC STUDY — 40 CORRIDOR COORDINATES & GOOGLE MAPS NAVIGATION LINKS (CAR & BIKE)"])
ws.cell(1, 1).font = title_font
ws.append(["Click the Car or Bike links to open Google Maps directly and verify origin, destination, and waypoints for each road corridor."])
ws.cell(2, 1).font = sub_font

headers = [
    "#", "Route ID", "Category", "Road Name / Corridor",
    "Origin Landmark", "Origin Coordinates (From)",
    "Destination Landmark", "Destination Coordinates (To)",
    "Via / Waypoint Coordinates (Lock Corridor)",
    "Car Navigation Link (Google Maps)",
    "Bike Navigation Link (Google Maps Two-Wheeler)"
]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(3, c)
    cell.fill = navy_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

row_num = 4
for r in routes:
    row_data = [
        r['index'],
        r['id'],
        r['section'],
        r['label'],
        r['originAddress'],
        r['from'],
        r['destAddress'],
        r['to'],
        r['viaCoord'] if r['viaCoord'] != 'None' else 'Direct Route (No Waypoints)',
        "Open Car Route",
        "Open Bike Route"
    ]
    ws.append(row_data)
    
    # Format cells
    for c in range(1, len(row_data) + 1):
        cell = ws.cell(row_num, c)
        cell.font = regular_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 6, 8] else "left", vertical="center")
        
    # Hyperlinks
    car_cell = ws.cell(row_num, 10)
    car_cell.hyperlink = r['carUrl']
    car_cell.font = link_font
    car_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    bike_cell = ws.cell(row_num, 11)
    bike_cell.hyperlink = r['bikeUrl']
    bike_cell.font = link_font
    bike_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    row_num += 1

# Autofit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        v = str(cell.value or '')
        if len(v) > max_len:
            max_len = len(v)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

for d in ['D:/Traffic', 'D:/Traffic/Bus/Data Collection 2']:
    p = os.path.join(d, 'Kolkata_Car_Bike_Route_Coordinates_and_Links.xlsx')
    wb.save(p)
    print(f"Saved: {p}")

# 2. Also generate an interactive HTML verification viewer
html_template = f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Kolkata Traffic Corridor Verification: Car & Bike</title>
<style>
  :root {{
    --primary: #1F4E78;
    --car: #C0392B;
    --bike: #27AE60;
    --bg: #F4F6F9;
    --card: #FFFFFF;
    --text: #2C3E50;
  }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; background: var(--bg); color: var(--text); padding: 25px; }}
  header {{ background: linear-gradient(135deg, #1F4E78 0%, #163753 100%); color: white; padding: 25px; border-radius: 10px; margin-bottom: 25px; }}
  header h1 {{ font-size: 24px; margin-bottom: 6px; }}
  header p {{ font-size: 13.5px; opacity: 0.9; }}
  
  .filter-bar {{ background: var(--card); padding: 15px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 2px 6px rgba(0,0,0,0.05); display: flex; gap: 15px; align-items: center; flex-wrap: wrap; }}
  input[type="text"] {{ padding: 8px 12px; border: 1px solid #BDC3C7; border-radius: 6px; font-size: 13px; width: 250px; }}
  .btn-filter {{ padding: 8px 14px; border: none; border-radius: 6px; background: #EAECEE; font-weight: bold; cursor: pointer; font-size: 12px; }}
  .btn-filter.active {{ background: var(--primary); color: white; }}

  table {{ width: 100%; border-collapse: collapse; background: var(--card); border-radius: 8px; overflow: hidden; box-shadow: 0 2px 8px rgba(0,0,0,0.06); font-size: 12.5px; }}
  th, td {{ padding: 10px 12px; text-align: left; border-bottom: 1px solid #ECF0F1; vertical-align: middle; }}
  th {{ background: var(--primary); color: white; font-weight: bold; }}
  tr:hover {{ background: #F8F9FA; }}
  
  .coord-pill {{ font-family: monospace; font-size: 11px; background: #EBF5FB; color: #1B4F72; padding: 3px 6px; border-radius: 4px; display: inline-block; margin-top: 3px; }}
  .via-pill {{ font-family: monospace; font-size: 10.5px; background: #FEF9E7; color: #7D6608; padding: 2px 5px; border-radius: 4px; display: block; margin-top: 2px; }}
  
  .btn-map {{ display: inline-block; padding: 5px 10px; border-radius: 5px; text-decoration: none; font-weight: bold; font-size: 11.5px; margin: 2px; text-align: center; }}
  .btn-car {{ background: #FDEDEC; color: var(--car); border: 1px solid var(--car); }}
  .btn-car:hover {{ background: var(--car); color: white; }}
  .btn-bike {{ background: #EAFAF1; color: var(--bike); border: 1px solid var(--bike); }}
  .btn-bike:hover {{ background: var(--bike); color: white; }}
</style>
</head>
<body>

<header>
  <h1>Kolkata 40 Corridor Verification Hub: Private Car & Motorcycle</h1>
  <p>Verify exact origin/destination coordinates, corridor waypoints, and open live Google Maps routes for verification prior to resuming automated data collection.</p>
</header>

<div class="filter-bar">
  <input type="text" id="searchInput" placeholder="Search by road name or route ID..." onkeyup="filterTable()">
  <button class="btn-filter active" onclick="setFilter('all')">All Corridors (40)</button>
  <button class="btn-filter" onclick="setFilter('Major Road')">Major Roads (22)</button>
  <button class="btn-filter" onclick="setFilter('Local Connector Segment')">Local Connectors (18)</button>
</div>

<table id="routesTable">
  <thead>
    <tr>
      <th style="width: 50px;">ID</th>
      <th style="width: 140px;">Road Name</th>
      <th>Category</th>
      <th>Origin (Start)</th>
      <th>Destination (End)</th>
      <th>Waypoints (Via)</th>
      <th style="width: 170px;">Google Maps Links</th>
    </tr>
  </thead>
  <tbody>
'''

for r in routes:
    via_html = '<span style="color:#7F8C8D;font-style:italic;">Direct</span>'
    if r['viaCoord'] != 'None':
        parts = r['viaCoord'].split('|')
        via_html = "".join([f'<span class="via-pill">📍 {p.strip()}</span>' for p in parts])
        
    html_template += f'''    <tr data-cat="{r['section']}">
      <td><strong>{r['id']}</strong></td>
      <td><strong>{r['label']}</strong></td>
      <td><span style="font-size:11px;color:#5D6D7E;">{r['section']}</span></td>
      <td>
        <div style="font-weight:600;">{r['originAddress']}</div>
        <div class="coord-pill">{r['from']}</div>
      </td>
      <td>
        <div style="font-weight:600;">{r['destAddress']}</div>
        <div class="coord-pill">{r['to']}</div>
      </td>
      <td>{via_html}</td>
      <td>
        <a class="btn-map btn-car" href="{r['carUrl']}" target="_blank">🚗 Open Car</a>
        <a class="btn-map btn-bike" href="{r['bikeUrl']}" target="_blank">🏍️ Open Bike</a>
      </td>
    </tr>
'''

html_template += '''  </tbody>
</table>

<script>
  let currentFilter = 'all';

  function setFilter(cat) {
    currentFilter = cat;
    document.querySelectorAll('.btn-filter').forEach(b => b.classList.remove('active'));
    event.target.classList.add('active');
    filterTable();
  }

  function filterTable() {
    const q = document.getElementById('searchInput').value.toLowerCase();
    const rows = document.querySelectorAll('#routesTable tbody tr');
    rows.forEach(r => {
      const cat = r.getAttribute('data-cat');
      const text = r.innerText.toLowerCase();
      const matchCat = (currentFilter === 'all' || cat === currentFilter);
      const matchText = text.includes(q);
      r.style.display = (matchCat && matchText) ? '' : 'none';
    });
  }
</script>

</body>
</html>
'''

for d in ['D:/Traffic', 'D:/Traffic/Bus/Data Collection 2']:
    hp = os.path.join(d, 'Route_Verification_Car_Bike.html')
    with open(hp, 'w', encoding='utf-8') as f:
        f.write(html_template)
    print(f"Saved: {hp}")

